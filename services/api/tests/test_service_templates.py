import io
from datetime import date
from decimal import Decimal
from uuid import uuid4

import pytest
from openpyxl import Workbook

from app.domain.constants import ConstructionMeasurementStatus, ConstructionProjectStatus
from app.domain.exceptions import (
    ConstructionDuplicateCodeError,
    ConstructionInvalidValueError,
    ConstructionResourceInUseError,
)
from app.domain.services import ConstructionProjectService
from app.domain.services.construction_service_template_parser import (
    SAMPLE_ITEMS,
    build_service_template_example,
    parse_service_template_spreadsheet,
)
from app.infrastructure.database.models import (
    ConstructionMeasurement,
    ConstructionProject,
    ConstructionUnit,
)
from app.schemas.construction import (
    ConstructionMeasurementItemCreate,
    ConstructionServiceTemplateCreate,
    ConstructionServiceTemplateItemInput,
    ConstructionServiceTemplateReplace,
    ConstructionServiceTemplateResponse,
    ConstructionServiceTemplateSectionInput,
    ConstructionServiceTemplateUpdate,
)
from tests.test_construction_project_service import FakeConstructionRepository


def build_caixa_spreadsheet(*, service_name, items, service_column=3, trailing_noise=True):
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.cell(row=1, column=service_column, value=f"Serviço: {service_name}")
    worksheet.cell(row=2, column=1, value="FICHA DE VERIFICAÇÃO DE SERVIÇO")
    worksheet.cell(row=3, column=1, value="Obra:")
    worksheet.cell(row=4, column=1, value="ITEM")
    worksheet.cell(row=4, column=3, value="MÉTODO DE VERIFICAÇÃO")
    for offset, (description, method) in enumerate(items):
        worksheet.cell(row=5 + offset, column=1, value=description)
        worksheet.cell(row=5 + offset, column=3, value=method)

    if trailing_noise:
        worksheet.cell(row=5 + len(items) + 1, column=1, value="LINHA APOS O BRANCO")

    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def test_parser_reads_service_name_and_inspection_items() -> None:
    content = build_caixa_spreadsheet(
        service_name="Alvenaria de vedação",
        items=[
            ("Locação da alvenaria", "Trena e esquadro"),
            ("Prumo das paredes", "Régua de 2m"),
        ],
    )

    parsed = parse_service_template_spreadsheet(file_name="fvs.xlsx", content=content)

    assert parsed.name == "ALVENARIA DE VEDAÇÃO"
    assert [(item.sequence_number, item.description, item.verification_method) for item in parsed.items] == [
        (1, "LOCAÇÃO DA ALVENARIA", "TRENA E ESQUADRO"),
        (2, "PRUMO DAS PAREDES", "RÉGUA DE 2M"),
    ]


def test_parser_reads_service_name_from_the_fallback_column() -> None:
    content = build_caixa_spreadsheet(
        service_name="Cobertura",
        items=[("Alinhamento das telhas", "Inspeção visual")],
        service_column=4,
    )

    parsed = parse_service_template_spreadsheet(file_name="fvs.xlsx", content=content)

    assert parsed.name == "COBERTURA"


def test_parser_stops_on_the_first_blank_row() -> None:
    content = build_caixa_spreadsheet(
        service_name="Pintura",
        items=[("Demão única", "Visual")],
        trailing_noise=True,
    )

    parsed = parse_service_template_spreadsheet(file_name="fvs.xlsx", content=content)

    assert len(parsed.items) == 1


def test_parser_ignores_duplicated_items() -> None:
    content = build_caixa_spreadsheet(
        service_name="Esquadrias",
        items=[
            ("Vedação do batente", "Visual"),
            ("Vedação do batente", "Outro método"),
            ("Prumo do batente", "Régua"),
        ],
    )

    parsed = parse_service_template_spreadsheet(file_name="fvs.xlsx", content=content)

    assert [item.description for item in parsed.items] == ["VEDAÇÃO DO BATENTE", "PRUMO DO BATENTE"]


def test_parser_rejects_spreadsheet_without_the_service_header() -> None:
    workbook = Workbook()
    workbook.active.cell(row=1, column=3, value="Planilha qualquer")
    buffer = io.BytesIO()
    workbook.save(buffer)

    with pytest.raises(ConstructionInvalidValueError) as error:
        parse_service_template_spreadsheet(file_name="qualquer.xlsx", content=buffer.getvalue())

    assert error.value.error_code == "CONSTRUCTION_TEMPLATE_SERVICE_NOT_FOUND"


def test_parser_rejects_spreadsheet_without_inspection_items() -> None:
    workbook = Workbook()
    workbook.active.cell(row=1, column=3, value="Serviço: Sem itens")
    buffer = io.BytesIO()
    workbook.save(buffer)

    with pytest.raises(ConstructionInvalidValueError) as error:
        parse_service_template_spreadsheet(file_name="vazia.xlsx", content=buffer.getvalue())

    assert error.value.error_code == "CONSTRUCTION_TEMPLATE_WITHOUT_ITEMS"


def test_parser_rejects_legacy_xls_format() -> None:
    with pytest.raises(ConstructionInvalidValueError) as error:
        parse_service_template_spreadsheet(file_name="fvs.xls", content=b"qualquer")

    assert error.value.error_code == "CONSTRUCTION_TEMPLATE_UNSUPPORTED_FORMAT"


def build_service(*, company_id=None):
    return ConstructionProjectService(repository=FakeConstructionRepository()), company_id or uuid4()


@pytest.mark.asyncio
async def test_import_creates_template_with_its_inspection_items() -> None:
    service, company_id = build_service()
    content = build_caixa_spreadsheet(
        service_name="Alvenaria",
        items=[("Locação", "Trena"), ("Prumo", "Régua")],
    )

    results = await service.import_service_templates(
        company_id=company_id,
        files=[("alvenaria.xlsx", content)],
    )

    assert results[0]["status"] == "created"
    assert results[0]["items_count"] == 2

    templates = await service.list_service_templates(company_id=company_id)

    assert len(templates) == 1
    assert templates[0].name == "ALVENARIA"
    assert [item.description for item in templates[0].items] == ["LOCAÇÃO", "PRUMO"]


@pytest.mark.asyncio
async def test_reimporting_the_same_service_is_skipped() -> None:
    service, company_id = build_service()
    content = build_caixa_spreadsheet(service_name="Alvenaria", items=[("Locação", "Trena")])

    await service.import_service_templates(company_id=company_id, files=[("alvenaria.xlsx", content)])
    results = await service.import_service_templates(company_id=company_id, files=[("alvenaria.xlsx", content)])

    assert results[0]["status"] == "skipped"

    templates = await service.list_service_templates(company_id=company_id)

    assert len(templates) == 1
    assert len(templates[0].items) == 1


@pytest.mark.asyncio
async def test_import_reports_failure_per_file_without_losing_the_others() -> None:
    service, company_id = build_service()
    valid_content = build_caixa_spreadsheet(service_name="Cobertura", items=[("Telhas", "Visual")])

    results = await service.import_service_templates(
        company_id=company_id,
        files=[("quebrada.xls", b"nao e xlsx"), ("cobertura.xlsx", valid_content)],
    )

    assert results[0]["status"] == "failed"
    assert results[1]["status"] == "created"
    assert len(await service.list_service_templates(company_id=company_id)) == 1


@pytest.mark.asyncio
async def test_import_without_file_is_refused() -> None:
    service, company_id = build_service()

    with pytest.raises(ConstructionInvalidValueError) as error:
        await service.import_service_templates(company_id=company_id, files=[])

    assert error.value.error_code == "CONSTRUCTION_TEMPLATE_NO_FILE"


async def build_measurement_scenario(*, service, company_id):
    repository = service.repository
    project = ConstructionProject(
        id=uuid4(),
        company_id=company_id,
        code="OBRA-070",
        name="Template project",
        status=ConstructionProjectStatus.ACTIVE,
        analytic_cost_center_id=uuid4(),
    )
    unit = ConstructionUnit(
        id=uuid4(),
        company_id=company_id,
        project_id=project.id,
        code="E-501",
        unit_type="house",
        analytic_cost_center_id=uuid4(),
        status="available",
    )
    measurement = ConstructionMeasurement(
        id=uuid4(),
        company_id=company_id,
        project_id=project.id,
        unit_id=unit.id,
        schedule_phase_id=uuid4(),
        code="MED-070",
        sequence_number=1,
        gross_amount=Decimal("1000.00"),
        retentions_amount=Decimal("0"),
        net_amount=Decimal("1000.00"),
        measured_amount=Decimal("1000.00"),
        due_date=date(2026, 10, 10),
        status=ConstructionMeasurementStatus.DRAFT,
    )
    repository.projects[(company_id, project.id)] = project
    repository.units[(company_id, unit.id)] = unit
    repository.measurements[(company_id, measurement.id)] = measurement
    return measurement


@pytest.mark.asyncio
async def test_measurement_item_from_template_copies_the_inspection_items() -> None:
    service, company_id = build_service()
    content = build_caixa_spreadsheet(
        service_name="Alvenaria",
        items=[("Locação", "Trena"), ("Prumo", "Régua"), ("Encunhamento", "Visual")],
    )
    await service.import_service_templates(company_id=company_id, files=[("alvenaria.xlsx", content)])
    template = (await service.list_service_templates(company_id=company_id))[0]
    measurement = await build_measurement_scenario(service=service, company_id=company_id)

    item = await service.create_measurement_item(
        company_id=company_id,
        measurement_id=measurement.id,
        request=ConstructionMeasurementItemCreate(
            service_template_id=template.id,
            amount=Decimal("1500.00"),
        ),
    )

    inspections = await service.repository.list_measurement_item_inspections(
        company_id=company_id,
        measurement_item_id=item.id,
    )

    assert item.description == "ALVENARIA"
    assert item.service_template_id == template.id
    assert [
        (inspection.sequence_number, inspection.description, inspection.verification_method)
        for inspection in inspections
    ] == [
        (1, "LOCAÇÃO", "TRENA"),
        (2, "PRUMO", "RÉGUA"),
        (3, "ENCUNHAMENTO", "VISUAL"),
    ]


@pytest.mark.asyncio
async def test_the_same_service_cannot_be_measured_twice_in_one_measurement() -> None:
    service, company_id = build_service()
    content = build_caixa_spreadsheet(service_name="Alvenaria", items=[("Locação", "Trena")])
    await service.import_service_templates(company_id=company_id, files=[("alvenaria.xlsx", content)])
    template = (await service.list_service_templates(company_id=company_id))[0]
    measurement = await build_measurement_scenario(service=service, company_id=company_id)

    await service.create_measurement_item(
        company_id=company_id,
        measurement_id=measurement.id,
        request=ConstructionMeasurementItemCreate(service_template_id=template.id, amount=Decimal("100.00")),
    )

    with pytest.raises(ConstructionInvalidValueError) as error:
        await service.create_measurement_item(
            company_id=company_id,
            measurement_id=measurement.id,
            request=ConstructionMeasurementItemCreate(service_template_id=template.id, amount=Decimal("200.00")),
        )

    assert error.value.error_code == "CONSTRUCTION_MEASUREMENT_ITEM_DUPLICATE_SERVICE"


@pytest.mark.asyncio
async def test_item_without_description_or_template_is_refused() -> None:
    service, company_id = build_service()
    measurement = await build_measurement_scenario(service=service, company_id=company_id)

    with pytest.raises(ConstructionInvalidValueError) as error:
        await service.create_measurement_item(
            company_id=company_id,
            measurement_id=measurement.id,
            request=ConstructionMeasurementItemCreate(amount=Decimal("100.00")),
        )

    assert error.value.error_code == "CONSTRUCTION_MEASUREMENT_ITEM_DESCRIPTION_REQUIRED"


@pytest.mark.asyncio
async def test_future_start_date_is_refused() -> None:
    service, company_id = build_service()
    measurement = await build_measurement_scenario(service=service, company_id=company_id)

    with pytest.raises(ConstructionInvalidValueError) as error:
        await service.create_measurement_item(
            company_id=company_id,
            measurement_id=measurement.id,
            request=ConstructionMeasurementItemCreate(
                description="Servico futuro",
                amount=Decimal("100.00"),
                start_date=date(2099, 1, 1),
            ),
        )

    assert error.value.error_code == "CONSTRUCTION_INVALID_PERIOD"


@pytest.mark.asyncio
async def test_end_date_is_blocked_while_an_inspection_is_pending() -> None:
    service, company_id = build_service()
    content = build_caixa_spreadsheet(service_name="Alvenaria", items=[("Locação", "Trena")])
    await service.import_service_templates(company_id=company_id, files=[("alvenaria.xlsx", content)])
    template = (await service.list_service_templates(company_id=company_id))[0]
    measurement = await build_measurement_scenario(service=service, company_id=company_id)

    with pytest.raises(ConstructionInvalidValueError) as error:
        await service.create_measurement_item(
            company_id=company_id,
            measurement_id=measurement.id,
            request=ConstructionMeasurementItemCreate(
                service_template_id=template.id,
                amount=Decimal("100.00"),
                start_date=date(2026, 1, 10),
                end_date=date(2026, 1, 20),
            ),
        )

    assert error.value.error_code == "CONSTRUCTION_MEASUREMENT_ITEM_END_DATE_BLOCKED"


@pytest.mark.asyncio
async def test_renaming_a_template_to_an_existing_name_is_refused() -> None:
    service, company_id = build_service()
    first = build_caixa_spreadsheet(service_name="Alvenaria", items=[("Locação", "Trena")])
    second = build_caixa_spreadsheet(service_name="Cobertura", items=[("Telhas", "Visual")])
    await service.import_service_templates(
        company_id=company_id,
        files=[("alvenaria.xlsx", first), ("cobertura.xlsx", second)],
    )
    templates = await service.list_service_templates(company_id=company_id)
    cobertura = next(template for template in templates if template.name == "COBERTURA")

    with pytest.raises(Exception) as error:
        await service.update_service_template(
            company_id=company_id,
            service_template_id=cobertura.id,
            request=ConstructionServiceTemplateUpdate(name="Alvenaria"),
        )

    assert "ALVENARIA" in str(error.value)


def test_example_spreadsheet_is_readable_by_the_parser_that_ships_with_it() -> None:
    content = build_service_template_example()

    parsed = parse_service_template_spreadsheet(file_name="modelo-fvs.xlsx", content=content)

    assert parsed.name == "ATERRO / COMPACTAÇÃO DE ATERRO"
    assert [section.name for section in parsed.sections] == ["GERAL"]
    assert len(parsed.items) == len(SAMPLE_ITEMS)
    assert parsed.items[0].description == SAMPLE_ITEMS[0][0].strip().upper()
    assert parsed.items[0].verification_method == SAMPLE_ITEMS[0][1].strip().upper()


@pytest.mark.asyncio
async def test_example_spreadsheet_can_be_imported_as_a_real_service() -> None:
    service, company_id = build_service()

    results = await service.import_service_templates(
        company_id=company_id,
        files=[("modelo-fvs.xlsx", build_service_template_example())],
    )

    assert results[0]["status"] == "created"
    assert results[0]["items_count"] == len(SAMPLE_ITEMS)


def build_section_input(*, name, items):
    return ConstructionServiceTemplateSectionInput(
        name=name,
        items=[
            ConstructionServiceTemplateItemInput(
                description=description,
                verification_method=method,
                requires_comment=requires_comment,
                requires_photo=requires_photo,
            )
            for description, method, requires_comment, requires_photo in items
        ],
    )


def test_parser_puts_every_item_in_a_single_default_section() -> None:
    content = build_caixa_spreadsheet(
        service_name="Aterro",
        items=[("Terreno limpo", "Visual"), ("Camadas", "Trena")],
    )

    parsed = parse_service_template_spreadsheet(file_name="aterro.xlsx", content=content)

    assert [(section.sequence_number, section.name) for section in parsed.sections] == [(1, "GERAL")]
    assert [(item.sequence_number, item.description) for item in parsed.items] == [
        (1, "TERRENO LIMPO"),
        (2, "CAMADAS"),
    ]


@pytest.mark.asyncio
async def test_import_creates_one_default_section_without_required_flags() -> None:
    service, company_id = build_service()
    content = build_caixa_spreadsheet(service_name="Aterro", items=[("Terreno limpo", "Visual")])

    results = await service.import_service_templates(company_id=company_id, files=[("aterro.xlsx", content)])

    assert results[0]["sections_count"] == 1
    template = (await service.list_service_templates(company_id=company_id))[0]
    assert [section.name for section in template.sections] == ["GERAL"]
    assert template.items[0].requires_comment is False
    assert template.items[0].requires_photo is False


@pytest.mark.asyncio
async def test_created_template_numbers_items_per_section() -> None:
    service, company_id = build_service()

    template = await service.create_service_template(
        company_id=company_id,
        request=ConstructionServiceTemplateCreate(
            name="Acabamento",
            sections=[
                build_section_input(
                    name="Materiais de acabamento",
                    items=[("Pisos ceramicos", "Visual", False, False), ("Tintas", "Visual", True, True)],
                ),
                build_section_input(
                    name="Materiais hidraulicos",
                    items=[("Instalacoes sanitarias", "Visual", True, False)],
                ),
            ],
        ),
    )

    assert [(section.sequence_number, section.name) for section in template.sections] == [
        (1, "MATERIAIS DE ACABAMENTO"),
        (2, "MATERIAIS HIDRAULICOS"),
    ]
    assert [(item.sequence_number, item.description) for item in template.sections[0].items] == [
        (1, "PISOS CERAMICOS"),
        (2, "TINTAS"),
    ]
    assert [(item.sequence_number, item.description) for item in template.sections[1].items] == [
        (1, "INSTALACOES SANITARIAS"),
    ]
    assert template.sections[0].items[1].requires_comment is True
    assert template.sections[0].items[1].requires_photo is True


@pytest.mark.asyncio
async def test_flat_items_follow_the_section_order_not_the_sequence_number() -> None:
    service, company_id = build_service()

    template = await service.create_service_template(
        company_id=company_id,
        request=ConstructionServiceTemplateCreate(
            name="Acabamento",
            sections=[
                build_section_input(
                    name="Primeira",
                    items=[("A1", "Visual", False, False), ("A2", "Visual", False, False)],
                ),
                build_section_input(name="Segunda", items=[("B1", "Visual", False, False)]),
            ],
        ),
    )

    response = ConstructionServiceTemplateResponse.from_model(template)

    assert [item.description for item in response.items] == ["A1", "A2", "B1"]


@pytest.mark.asyncio
async def test_creating_a_template_with_an_existing_name_is_refused() -> None:
    service, company_id = build_service()
    await service.create_service_template(
        company_id=company_id,
        request=ConstructionServiceTemplateCreate(name="Aterro", sections=[]),
    )

    with pytest.raises(ConstructionDuplicateCodeError):
        await service.create_service_template(
            company_id=company_id,
            request=ConstructionServiceTemplateCreate(name="aterro", sections=[]),
        )


@pytest.mark.asyncio
async def test_repeated_section_name_is_refused() -> None:
    service, company_id = build_service()

    with pytest.raises(ConstructionInvalidValueError) as error:
        await service.create_service_template(
            company_id=company_id,
            request=ConstructionServiceTemplateCreate(
                name="Acabamento",
                sections=[
                    build_section_input(name="Materiais", items=[("A1", "Visual", False, False)]),
                    build_section_input(name="materiais", items=[("B1", "Visual", False, False)]),
                ],
            ),
        )

    assert error.value.error_code == "CONSTRUCTION_TEMPLATE_DUPLICATE_SECTION"


@pytest.mark.asyncio
async def test_repeated_item_inside_a_section_is_refused_but_across_sections_is_allowed() -> None:
    service, company_id = build_service()

    with pytest.raises(ConstructionInvalidValueError) as error:
        await service.create_service_template(
            company_id=company_id,
            request=ConstructionServiceTemplateCreate(
                name="Acabamento",
                sections=[
                    build_section_input(
                        name="Materiais",
                        items=[("Pisos", "Visual", False, False), ("pisos", "Trena", False, False)],
                    ),
                ],
            ),
        )

    assert error.value.error_code == "CONSTRUCTION_TEMPLATE_DUPLICATE_ITEM"

    template = await service.create_service_template(
        company_id=company_id,
        request=ConstructionServiceTemplateCreate(
            name="Cobertura",
            sections=[
                build_section_input(name="Primeira", items=[("Pisos", "Visual", False, False)]),
                build_section_input(name="Segunda", items=[("Pisos", "Trena", False, False)]),
            ],
        ),
    )

    assert [len(section.items) for section in template.sections] == [1, 1]


@pytest.mark.asyncio
async def test_replacing_a_template_drops_the_sections_that_left() -> None:
    service, company_id = build_service()
    template = await service.create_service_template(
        company_id=company_id,
        request=ConstructionServiceTemplateCreate(
            name="Acabamento",
            sections=[
                build_section_input(name="Primeira", items=[("A1", "Visual", False, False)]),
                build_section_input(name="Segunda", items=[("B1", "Visual", False, False)]),
            ],
        ),
    )

    replaced = await service.replace_service_template(
        company_id=company_id,
        service_template_id=template.id,
        request=ConstructionServiceTemplateReplace(
            name="Acabamento revisado",
            sections=[build_section_input(name="Unica", items=[("C1", "Visual", False, False)])],
        ),
    )

    assert replaced.name == "ACABAMENTO REVISADO"
    assert [section.name for section in replaced.sections] == ["UNICA"]
    assert [item.description for item in replaced.items] == ["C1"]


@pytest.mark.asyncio
async def test_replacing_a_template_in_use_keeps_the_inspections_already_copied() -> None:
    service, company_id = build_service()
    template = await service.create_service_template(
        company_id=company_id,
        request=ConstructionServiceTemplateCreate(
            name="Alvenaria",
            sections=[build_section_input(name="Geral", items=[("Locacao", "Trena", False, False)])],
        ),
    )
    measurement = await build_measurement_scenario(service=service, company_id=company_id)
    item = await service.create_measurement_item(
        company_id=company_id,
        measurement_id=measurement.id,
        request=ConstructionMeasurementItemCreate(amount=Decimal("100.00"), service_template_id=template.id),
        actor_user_id=uuid4(),
    )

    await service.replace_service_template(
        company_id=company_id,
        service_template_id=template.id,
        request=ConstructionServiceTemplateReplace(
            name="Alvenaria",
            sections=[build_section_input(name="Outra", items=[("Prumo", "Regua", False, False)])],
        ),
    )

    inspections = await service.repository.list_measurement_item_inspections(
        company_id=company_id,
        measurement_item_id=item.id,
    )
    assert [inspection.description for inspection in inspections] == ["LOCACAO"]


@pytest.mark.asyncio
async def test_deleting_an_unused_template_removes_it_from_the_catalog() -> None:
    service, company_id = build_service()
    template = await service.create_service_template(
        company_id=company_id,
        request=ConstructionServiceTemplateCreate(
            name="Aterro",
            sections=[build_section_input(name="Geral", items=[("Terreno", "Visual", False, False)])],
        ),
    )

    await service.delete_service_template(company_id=company_id, service_template_id=template.id)

    assert await service.list_service_templates(company_id=company_id) == []


@pytest.mark.asyncio
async def test_deleting_a_template_used_by_a_measurement_is_refused() -> None:
    service, company_id = build_service()
    template = await service.create_service_template(
        company_id=company_id,
        request=ConstructionServiceTemplateCreate(
            name="Alvenaria",
            sections=[build_section_input(name="Geral", items=[("Locacao", "Trena", False, False)])],
        ),
    )
    measurement = await build_measurement_scenario(service=service, company_id=company_id)
    await service.create_measurement_item(
        company_id=company_id,
        measurement_id=measurement.id,
        request=ConstructionMeasurementItemCreate(amount=Decimal("100.00"), service_template_id=template.id),
        actor_user_id=uuid4(),
    )

    with pytest.raises(ConstructionResourceInUseError) as error:
        await service.delete_service_template(company_id=company_id, service_template_id=template.id)

    assert error.value.error_code == "CONSTRUCTION_TEMPLATE_IN_USE"
    assert error.value.status_code == 409
    assert "Desative" in error.value.message
    assert len(await service.list_service_templates(company_id=company_id)) == 1


@pytest.mark.asyncio
async def test_inspections_copied_from_a_multi_section_template_are_numbered_continuously() -> None:
    service, company_id = build_service()
    template = await service.create_service_template(
        company_id=company_id,
        request=ConstructionServiceTemplateCreate(
            name="Acabamento",
            sections=[
                build_section_input(
                    name="Materiais",
                    items=[("Pisos", "Visual", False, False), ("Tintas", "Visual", True, True)],
                ),
                build_section_input(name="Hidraulicos", items=[("Sanitarias", "Visual", True, False)]),
            ],
        ),
    )
    measurement = await build_measurement_scenario(service=service, company_id=company_id)

    item = await service.create_measurement_item(
        company_id=company_id,
        measurement_id=measurement.id,
        request=ConstructionMeasurementItemCreate(amount=Decimal("100.00"), service_template_id=template.id),
        actor_user_id=uuid4(),
    )

    inspections = await service.repository.list_measurement_item_inspections(
        company_id=company_id,
        measurement_item_id=item.id,
    )
    assert [
        (inspection.sequence_number, inspection.section_name, inspection.description)
        for inspection in inspections
    ] == [
        (1, "MATERIAIS", "PISOS"),
        (2, "MATERIAIS", "TINTAS"),
        (3, "HIDRAULICOS", "SANITARIAS"),
    ]
    assert [inspection.requires_photo for inspection in inspections] == [False, True, False]
    assert [inspection.requires_comment for inspection in inspections] == [False, True, True]


class FakeErpPeopleClient:
    """So o suficiente para a auditoria: resolve o nome de uma pessoa."""

    def __init__(self, names: dict) -> None:
        self.names = names

    async def get_person_name(self, *, company_id, user_id, person_id):
        return self.names.get(person_id)


def build_service_with_people(names):
    return (
        ConstructionProjectService(
            repository=FakeConstructionRepository(),
            erp_client=FakeErpPeopleClient(names=names),
        ),
        uuid4(),
    )


@pytest.mark.asyncio
async def test_creating_a_service_records_who_created_it() -> None:
    service, company_id = build_service()
    actor_user_id = uuid4()

    template = await service.create_service_template(
        company_id=company_id,
        request=ConstructionServiceTemplateCreate(
            name="Aterro",
            sections=[build_section_input(name="Geral", items=[("Terreno", "Visual", False, False)])],
        ),
        actor_user_id=actor_user_id,
        actor_person_id=uuid4(),
    )

    assert template.created_by_user_id == actor_user_id
    assert template.updated_by_user_id == actor_user_id

    audits = await service.list_service_template_audits(
        company_id=company_id,
        service_template_id=template.id,
    )
    assert [audit.event for audit in audits] == ["created"]
    assert audits[0].actor_user_id == actor_user_id
    assert audits[0].service_template_name == "ATERRO"
    assert audits[0].summary == "1 seção, 1 item."


@pytest.mark.asyncio
async def test_revising_a_service_keeps_the_previous_author_in_the_trail() -> None:
    """O autor anterior so sobrevive na tabela de auditoria.

    As colunas do servico guardam apenas o ULTIMO autor -- e por isso que a
    coluna sozinha nao responderia quem escreveu a versao que a medicao usou.
    """
    service, company_id = build_service()
    author, reviewer = uuid4(), uuid4()

    template = await service.create_service_template(
        company_id=company_id,
        request=ConstructionServiceTemplateCreate(
            name="Aterro",
            sections=[build_section_input(name="Geral", items=[("Terreno", "Visual", False, False)])],
        ),
        actor_user_id=author,
    )
    revised = await service.replace_service_template(
        company_id=company_id,
        service_template_id=template.id,
        request=ConstructionServiceTemplateReplace(
            name="Aterro",
            sections=[
                build_section_input(
                    name="Geral",
                    items=[("Terreno", "Visual", False, False), ("Umidade", "Ensaio", True, False)],
                )
            ],
        ),
        actor_user_id=reviewer,
    )

    assert revised.created_by_user_id == author
    assert revised.updated_by_user_id == reviewer

    audits = await service.list_service_template_audits(
        company_id=company_id,
        service_template_id=template.id,
    )
    assert [audit.event for audit in audits] == ["replaced", "created"]
    assert audits[0].actor_user_id == reviewer
    assert audits[1].actor_user_id == author
    assert len(audits[0].snapshot["sections"][0]["items"]) == 2
    assert len(audits[1].snapshot["sections"][0]["items"]) == 1


@pytest.mark.asyncio
async def test_deactivating_and_reactivating_are_their_own_events() -> None:
    service, company_id = build_service()
    template = await service.create_service_template(
        company_id=company_id,
        request=ConstructionServiceTemplateCreate(
            name="Aterro",
            sections=[build_section_input(name="Geral", items=[("Terreno", "Visual", False, False)])],
        ),
        actor_user_id=uuid4(),
    )
    who_deactivated = uuid4()

    await service.update_service_template(
        company_id=company_id,
        service_template_id=template.id,
        request=ConstructionServiceTemplateUpdate(is_active=False),
        actor_user_id=who_deactivated,
    )
    await service.update_service_template(
        company_id=company_id,
        service_template_id=template.id,
        request=ConstructionServiceTemplateUpdate(is_active=True),
        actor_user_id=uuid4(),
    )

    audits = await service.list_service_template_audits(
        company_id=company_id,
        service_template_id=template.id,
    )
    assert [audit.event for audit in audits] == ["activated", "deactivated", "created"]
    assert audits[1].actor_user_id == who_deactivated


@pytest.mark.asyncio
async def test_renaming_a_service_says_the_old_name_in_the_trail() -> None:
    service, company_id = build_service()
    template = await service.create_service_template(
        company_id=company_id,
        request=ConstructionServiceTemplateCreate(name="Aterro", sections=[]),
        actor_user_id=uuid4(),
    )

    await service.update_service_template(
        company_id=company_id,
        service_template_id=template.id,
        request=ConstructionServiceTemplateUpdate(name="Aterro compactado"),
        actor_user_id=uuid4(),
    )

    audits = await service.list_service_template_audits(
        company_id=company_id,
        service_template_id=template.id,
    )
    assert audits[0].event == "updated"
    assert audits[0].summary == "Nome alterado de 'ATERRO' para 'ATERRO COMPACTADO'."
    # Denormalizado: o registro carrega o nome de ENTAO, nao o de agora.
    assert audits[1].service_template_name == "ATERRO"


@pytest.mark.asyncio
async def test_deleting_a_service_hides_it_but_keeps_its_history() -> None:
    service, company_id = build_service()
    who_deleted = uuid4()
    template = await service.create_service_template(
        company_id=company_id,
        request=ConstructionServiceTemplateCreate(
            name="Aterro",
            sections=[build_section_input(name="Geral", items=[("Terreno", "Visual", False, False)])],
        ),
        actor_user_id=uuid4(),
    )

    await service.delete_service_template(
        company_id=company_id,
        service_template_id=template.id,
        actor_user_id=who_deleted,
    )

    assert await service.list_service_templates(company_id=company_id) == []
    deleted = await service.list_service_templates(company_id=company_id, only_deleted=True)
    assert [item.name for item in deleted] == ["ATERRO"]
    assert deleted[0].deleted_by_user_id == who_deleted
    assert deleted[0].deleted_at is not None

    audits = await service.list_service_template_audits(
        company_id=company_id,
        service_template_id=template.id,
    )
    assert audits[0].event == "deleted"
    assert audits[0].actor_user_id == who_deleted
    # O snapshot e o da ficha VIVA: e ele que diz o que saiu do catalogo.
    assert len(audits[0].snapshot["sections"][0]["items"]) == 1


@pytest.mark.asyncio
async def test_a_deleted_service_releases_its_name() -> None:
    service, company_id = build_service()
    template = await service.create_service_template(
        company_id=company_id,
        request=ConstructionServiceTemplateCreate(name="Aterro", sections=[]),
        actor_user_id=uuid4(),
    )
    await service.delete_service_template(company_id=company_id, service_template_id=template.id)

    recreated = await service.create_service_template(
        company_id=company_id,
        request=ConstructionServiceTemplateCreate(name="Aterro", sections=[]),
        actor_user_id=uuid4(),
    )

    assert recreated.id != template.id
    # Cada um com o seu proprio historico: o excluido nao some do registro.
    audits = await service.list_service_template_audits(
        company_id=company_id,
        service_template_id=template.id,
    )
    assert [audit.event for audit in audits] == ["deleted", "created"]


@pytest.mark.asyncio
async def test_the_actor_name_is_frozen_at_the_moment_of_the_event() -> None:
    person_id = uuid4()
    service, company_id = build_service_with_people({person_id: "Maria de Souza"})

    template = await service.create_service_template(
        company_id=company_id,
        request=ConstructionServiceTemplateCreate(name="Aterro", sections=[]),
        actor_user_id=uuid4(),
        actor_person_id=person_id,
    )
    service.erp_client.names[person_id] = "Maria de Souza Lima"
    await service.update_service_template(
        company_id=company_id,
        service_template_id=template.id,
        request=ConstructionServiceTemplateUpdate(is_active=False),
        actor_user_id=uuid4(),
        actor_person_id=person_id,
    )

    audits = await service.list_service_template_audits(
        company_id=company_id,
        service_template_id=template.id,
    )
    assert audits[0].actor_name == "Maria de Souza Lima"
    assert audits[1].actor_name == "Maria de Souza"


@pytest.mark.asyncio
async def test_a_silent_erp_does_not_block_the_write() -> None:
    """Sem nome, mas com id: gravar tem de acontecer de qualquer jeito.

    Perder a alteracao para conseguir registrar quem a fez seria trocar o
    problema por outro maior.
    """
    service, company_id = build_service()

    template = await service.create_service_template(
        company_id=company_id,
        request=ConstructionServiceTemplateCreate(name="Aterro", sections=[]),
        actor_user_id=uuid4(),
        actor_person_id=uuid4(),
    )

    audits = await service.list_service_template_audits(
        company_id=company_id,
        service_template_id=template.id,
    )
    assert audits[0].actor_name is None
    assert audits[0].actor_user_id is not None


@pytest.mark.asyncio
async def test_importing_a_sheet_records_the_file_it_came_from() -> None:
    service, company_id = build_service()
    actor_user_id = uuid4()
    content = build_caixa_spreadsheet(
        service_name="Alvenaria",
        items=[("Locação", "Trena"), ("Prumo", "Régua")],
    )

    await service.import_service_templates(
        company_id=company_id,
        files=[("alvenaria.xlsx", content)],
        actor_user_id=actor_user_id,
    )

    template = (await service.list_service_templates(company_id=company_id))[0]
    assert template.created_by_user_id == actor_user_id

    audits = await service.list_service_template_audits(
        company_id=company_id,
        service_template_id=template.id,
    )
    assert audits[0].event == "created"
    assert audits[0].source == "import"
    assert audits[0].actor_user_id == actor_user_id
    assert "alvenaria.xlsx" in audits[0].summary
    assert len(audits[0].snapshot["sections"][0]["items"]) == 2
